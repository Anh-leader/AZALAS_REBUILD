"""Nạp dữ liệu nguồn Google Sheets (đã export ra .xlsx/.csv/.tsv).

Mỗi sheet được mô tả bằng alias cột (giống app.import_customers) để chịu được
tên header thay đổi nhẹ giữa các lần export. KHÔNG tự đoán: thiếu cột bắt buộc
thì báo lỗi rõ tên sheet + cột.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.import_customers import _header_key, _text


@dataclass
class SheetSpec:
    sheet: str
    aliases: dict[str, tuple[str, ...]]
    required: tuple[str, ...]


# Nguồn 02_DATA_SALES_CRM
SHEET_CUSTOMERS = SheetSpec(
    sheet="tbl_KhachHang",
    aliases={
        "customer_code": ("makhachhang", "makh", "makhach"),
        "name": ("tenkhachhang", "tenkhach", "khachhang", "ten"),
        "phone": ("sdt", "sodienthoai", "dienthoai", "phone"),
        "address": ("diachi", "address"),
        "customer_group": ("nhomkhach", "nhomkhachhang"),
        "zalo": ("zalo",),
        "area": ("khuvuc", "khuvucban"),
        "status": ("trangthai", "status"),
        "note": ("ghichu", "note"),
        "created": ("ngaytao", "created", "createdat"),
        "updated": ("capnhatcuoi", "updatedat"),
        "carrier_name": ("chanhxe", "chanh"),
        "carrier_phone": ("sdtchanhxe", "sdchanh"),
        "source_debt": ("sono", "nohientai", "congno"),
    },
    required=("customer_code", "name"),
)

SHEET_ORDERS = SheetSpec(
    sheet="tbl_DonHang",
    aliases={
        "order_no": ("mahoadon", "madh", "hoadon", "mahd"),
        "customer_code": ("makh", "makhachhang"),
        "picker": ("nguoisoan", "nguoinhan", "nguoixuly"),
        "status": ("trangthai", "trangthaidon"),
        "billing_status": ("trangthaihoadon", "trangthaihd"),
        "order_type": ("kieudon", "loaidon"),
        "created_date": ("ngaytao", "ngaytaodon"),
        "delivery_date": ("ngaygiao", "ngaygiao"),
        "total_amount": ("tongtien", "tongtienhang"),
        "paid_amount": ("dathanhtoan", "datra", "tiennhantruoc"),
        "invoice_debt": ("notronghoadon", "nohoadon"),
        "customer_note": ("ghichukhach", "ghichukh"),
        "staff_note": ("ghichunv", "ghichunhanvien"),
        "channel": ("kenhban", "kenh", "channel"),
        "warehouse": ("khoxuat", "kho"),
    },
    required=("order_no", "customer_code"),
)

SHEET_ORDER_ITEMS = SheetSpec(
    sheet="tbl_ChiTietDonHang",
    aliases={
        "order_no": ("mahoadon", "madh"),
        "product_name": ("tensp", "tensanpham", "sanpham"),
        "quantity": ("soluong", "sl"),
        "unit_price": ("dongia", "gia"),
        "line_total": ("thanhtien", "tongtien"),
        "sales_code": ("mabanhang", "mabh"),
        "internal_code": ("mahangnoibo", "manoibo", "mahh"),
        "color": ("mauchuan", "mau"),
        "size": ("size", "kichthuoc"),
        "warehouse": ("khoxuat", "kho"),
        "sale_type": ("kieuban", "loaiban"),
        "product_image": ("anhsanpham", "anh"),
        "note": ("ghichu", "note"),
    },
    required=("order_no", "product_name"),
)

SHEET_STATUS_HISTORY = SheetSpec(
    sheet="tbl_LichSuTrangThaiDon",
    aliases={
        "order_no": ("mahoadon", "madh"),
        "from_status": ("trangthaicu", "tutrangthai", "trangthaitu"),
        "to_status": ("trangthaimoi", "dentrangthai", "trangthaiden"),
        "reason": ("lydo", "ghichu", "noidung"),
        "occurred_at": ("thoigian", "ngay", "ngaytao", "occurredat"),
        "actor": ("nguoithuchien", "nguoidoi", "nguoisua"),
    },
    required=("order_no", "to_status"),
)

SHEET_MEDIA = SheetSpec(
    sheet="tbl_DonHang_AnhZalo",
    aliases={
        "ma_anh": ("maanh", "maanhzalo"),
        "order_no": ("mahoadon", "madh"),
        "customer_code": ("makh", "makhachhang"),
        "drive_link": ("linkanh", "drivelink", "link", "url"),
        "created_at": ("ngaytao", "thoigian", "ngay", "createdat"),
        "uploader": ("nguoiup", "nguoidang", "uploader"),
        "note": ("ghichu", "note"),
        "original_name": ("tenfile", "tenanh", "filename"),
    },
    required=("ma_anh", "order_no"),
)

# Nguồn 06_DATA_TAI_CHINH_CONG_NO
SHEET_DEBT_LEDGER = SheetSpec(
    sheet="tbl_NhatKyNo",
    aliases={
        "legacy_id": ("magiaodich", "maphieu", "id", "ma"),
        "customer_code": ("makh", "makhachhang"),
        "occurred_at": ("ngay", "thoigian", "ngaytao"),
        "entry_type": ("loai", "loaigiaodich", "kieugiaodich"),
        "amount": ("sotien", "sotien", "amount"),
        "direction": ("loaicongno", "dau", "direction", "congno"),
        "description": ("ghichu", "noidung", "lydo", "description"),
    },
    required=("customer_code", "amount"),
)

SHEET_PAYMENTS = SheetSpec(
    sheet="tbl_ThanhToanKhach",
    aliases={
        "legacy_id": ("maphieu", "magiaodich", "id", "ma"),
        "customer_code": ("makh", "makhachhang"),
        "paid_at": ("ngaythanhtoan", "ngay", "thoigian"),
        "amount": ("sotien", "sotienthanhtoan", "amount"),
        "method": ("hinhthuc", "phuongthuc", "method"),
        "note": ("ghichu", "note", "noidung"),
    },
    required=("customer_code", "amount"),
)

SHEET_DEBT_SUMMARY = SheetSpec(
    sheet="tbl_CongNo_TongHop",
    aliases={
        "customer_code": ("makh", "makhachhang"),
        "debt": ("sono", "nohientai", "congnohientai", "tongno"),
    },
    required=("customer_code", "debt"),
)

# Nguồn 03_DATA_KHO_THANH_PHAM
SHEET_VARIANTS = SheetSpec(
    sheet="tbl_Kho_BienThe",
    aliases={
        "internal_code": ("mahangnoibo", "manoibo"),
        "sales_code": ("mabanhang", "mabh"),
        "product_code": ("masp", "mahanghoa"),
        "color": ("mau", "mauchuan"),
        "size": ("size", "kichthuoc"),
        "display_name": ("tenhienthi", "tenbienthe"),
        "status": ("coban", "trangthai"),
    },
    required=("internal_code", "size"),
)

SHEET_INVENTORY = SheetSpec(
    sheet="tbl_Kho_Ton",
    aliases={
        "internal_code": ("mahangnoibo", "manoibo"),
        "sales_code": ("mabanhang", "mabh"),
        "warehouse": ("kho", "khoxuat"),
        "location": ("vitri", "location"),
        "on_hand": ("tonthucte", "ton", "tonkho"),
        "reserved": ("tongiu", "dangiu", "reserved"),
        "available": ("khadung", "available"),
        "pending_finish": ("chohoanthien",),
        "cut_separated": ("dacotri",),
        "defective": ("hangloi", "loi"),
    },
    required=("internal_code", "warehouse", "on_hand"),
)


PHASE1_SHEETS: dict[str, SheetSpec] = {
    spec.sheet: spec
    for spec in (
        SHEET_CUSTOMERS, SHEET_ORDERS, SHEET_ORDER_ITEMS, SHEET_STATUS_HISTORY,
        SHEET_MEDIA, SHEET_DEBT_LEDGER, SHEET_PAYMENTS, SHEET_DEBT_SUMMARY,
        SHEET_VARIANTS, SHEET_INVENTORY,
    )
}


@dataclass
class SourceRow:
    sheet: str
    row_number: int
    values: dict[str, str | None]


@dataclass
class SourceSheet:
    sheet: str
    rows: list[SourceRow]
    columns: dict[str, int] = field(default_factory=dict)


def _parse_rows(sheet: str, spec: SheetSpec, rows: Iterable[tuple[Any, ...]]) -> SourceSheet:
    iterator = iter(rows)
    try:
        headers = next(iterator)
    except StopIteration as exc:
        raise ValueError(f"{sheet} không có dòng tiêu đề") from exc
    normalized = {_header_key(value): index for index, value in enumerate(headers)}
    columns: dict[str, int] = {}
    for field_name, aliases in spec.aliases.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match is not None:
            columns[field_name] = match
    missing = [field_name for field_name in spec.required if field_name not in columns]
    if missing:
        raise ValueError(f"{sheet} thiếu cột bắt buộc: {', '.join(missing)}")

    parsed: list[SourceRow] = []
    for row_number, row in enumerate(iterator, start=2):
        if not any(_text(cell) for cell in row):
            continue
        values = {
            field_name: _text(row[index]) if index < len(row) else None
            for field_name, index in columns.items()
        }
        parsed.append(SourceRow(sheet=sheet, row_number=row_number, values=values))
    return SourceSheet(sheet=sheet, rows=parsed, columns=columns)


def load_sheet_from_file(path: Path, spec: SheetSpec) -> SourceSheet:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            matching = next(
                (name for name in workbook.sheetnames if _header_key(name) == _header_key(spec.sheet)),
                None,
            )
            if not matching:
                raise ValueError(
                    f"{path.name}: không tìm thấy sheet {spec.sheet}; hiện có: {', '.join(workbook.sheetnames)}"
                )
            return _parse_rows(spec.sheet, spec, workbook[matching].iter_rows(values_only=True))
        finally:
            workbook.close()
    if suffix in (".csv", ".tsv"):
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            return _parse_rows(spec.sheet, spec, (tuple(row) for row in csv.reader(source, delimiter=delimiter)))
    raise ValueError(f"Chỉ hỗ trợ .xlsx, .xlsm, .csv, .tsv: {path.name}")


def discover_sheets(source_dir: Path, wanted: Iterable[str]) -> dict[str, SourceSheet]:
    """Quét thư mục nguồn, tìm sheet theo tên chuẩn hóa, nạp sheet đầu tiên khớp."""
    if not source_dir.is_dir():
        raise ValueError(f"Thư mục nguồn không tồn tại: {source_dir}")
    files = sorted(
        [p for p in source_dir.iterdir() if p.suffix.lower() in (".xlsx", ".xlsm", ".csv", ".tsv")],
        key=lambda p: p.name,
    )
    if not files:
        raise ValueError(f"Thư mục nguồn trống: {source_dir}")

    wanted_keys = {_header_key(w) for w in wanted}
    result: dict[str, SourceSheet] = {}
    for sheet_name, spec in PHASE1_SHEETS.items():
        if _header_key(sheet_name) not in wanted_keys:
            continue
        found: SourceSheet | None = None
        last_error: Exception | None = None
        for path in files:
            try:
                candidate = load_sheet_from_file(path, spec)
                if candidate.rows:
                    found = candidate
                    break
            except ValueError as exc:
                last_error = exc
                continue
        if found is None:
            raise ValueError(f"Không nạp được sheet {sheet_name}: {last_error or 'không có file chứa sheet này'}")
        result[sheet_name] = found
    return result


def money_int(value: str | None, *, sheet: str, row: int, field: str) -> int:
    """Tiền VND luôn là số nguyên, không float."""
    if value is None:
        return 0
    text = value.replace(" ", "").replace(",", "").replace(".", "")
    if text in ("", "-"):
        return 0
    try:
        return int(round(float(text)))
    except ValueError as exc:
        raise ValueError(f"{sheet} dòng {row}: {field} không phải số tiền hợp lệ: {value}") from exc


def decimal_qty(value: str | None, *, sheet: str, row: int, field: str):
    from decimal import Decimal, InvalidOperation
    if value is None:
        return Decimal(0)
    try:
        result = Decimal(value.replace(" ", "").replace(",", "."))
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError(f"{sheet} dòng {row}: {field} không phải số: {value}") from exc
    if not result.is_finite() or result < 0:
        raise ValueError(f"{sheet} dòng {row}: {field} phải >= 0")
    return result


def parse_date(value: str | None) -> str | None:
    """Chuẩn hóa ngày sang ISO YYYY-MM-DD; None nếu không parse được."""
    if not value:
        return None
    from datetime import datetime
    text = value.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None
