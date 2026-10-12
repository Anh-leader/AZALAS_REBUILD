"""Import `tbl_KhachHang` from an exported 02_DATA_SALES_CRM workbook.

Safety rules:
- `MaKH` is the sole upsert identity; customer names are never keys.
- Existing customers are skipped, never updated.
- Phone is copied only when the source `SDT` cell is non-empty.
- No phone, customer code or other real-data field is synthesized.
"""

from __future__ import annotations

import argparse
import csv
import math
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .core.db import SessionLocal
from .models import Customer


SHEET_NAME = "tbl_KhachHang"
HEADER_ALIASES = {
    "customer_code": ("makh", "makhachhang", "makhach"),
    "name": ("tenkhach", "tenkhachhang", "khachhang", "ten"),
    "phone": ("sdt", "sodienthoai", "dienthoai", "phone"),
    "address": ("diachi", "address"),
    "customer_group": ("nhomkhach", "nhomkhachhang", "customergroup"),
    "status": ("trangthai", "status"),
}


@dataclass(frozen=True)
class CustomerImportRow:
    row_number: int
    customer_code: str
    name: str
    phone: str | None
    address: str | None
    customer_group: str | None
    source_status: str | None


@dataclass(frozen=True)
class ImportSummary:
    source_rows: int
    imported: int
    skipped_existing: int


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _header_key(value: Any) -> str:
    text = unicodedata.normalize("NFD", _text(value) or "")
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return "".join(char for char in text.lower().replace("đ", "d") if char.isalnum())


def _column_map(headers: Iterable[Any]) -> dict[str, int]:
    normalized = {_header_key(value): index for index, value in enumerate(headers)}
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match is not None:
            result[field] = match
    missing = [field for field in ("customer_code", "name") if field not in result]
    if missing:
        raise ValueError(f"Thiếu cột bắt buộc trong {SHEET_NAME}: {', '.join(missing)}")
    return result


def _parse_rows(rows: Iterable[tuple[Any, ...]]) -> list[CustomerImportRow]:
    iterator = iter(rows)
    try:
        headers = next(iterator)
    except StopIteration as exc:
        raise ValueError(f"{SHEET_NAME} không có dòng tiêu đề") from exc
    columns = _column_map(headers)
    parsed: list[CustomerImportRow] = []
    seen_codes: set[str] = set()

    def value(row: tuple[Any, ...], field: str) -> str | None:
        index = columns.get(field)
        return _text(row[index]) if index is not None and index < len(row) else None

    for row_number, row in enumerate(iterator, start=2):
        if not any(_text(cell) for cell in row):
            continue
        code = value(row, "customer_code")
        name = value(row, "name")
        if not code or not name:
            raise ValueError(f"Dòng {row_number}: MaKH và TenKhach là bắt buộc")
        if code in seen_codes:
            raise ValueError(f"Dòng {row_number}: MaKH bị trùng trong file nguồn: {code}")
        if len(code) > 40 or len(name) > 300:
            raise ValueError(f"Dòng {row_number}: MaKH hoặc TenKhach vượt độ dài schema")
        phone = value(row, "phone")
        if phone and len(phone) > 30:
            raise ValueError(f"Dòng {row_number}: SDT vượt quá 30 ký tự")
        seen_codes.add(code)
        parsed.append(CustomerImportRow(
            row_number=row_number,
            customer_code=code,
            name=name,
            phone=phone,  # Giữ NULL khi ô SDT trống; không tự sinh/điền số.
            address=value(row, "address"),
            customer_group=value(row, "customer_group"),
            source_status=value(row, "status"),
        ))
    return parsed


def load_customer_rows(path: Path, sheet_name: str = SHEET_NAME) -> list[CustomerImportRow]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return _load_workbook_rows(workbook, sheet_name)
        finally:
            workbook.close()
    if suffix in (".csv", ".tsv"):
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            return _parse_rows(tuple(row) for row in csv.reader(source, delimiter=delimiter))
    raise ValueError("Chỉ hỗ trợ file .xlsx, .xlsm, .csv hoặc .tsv")


def _load_workbook_rows(workbook: Any, sheet_name: str = SHEET_NAME) -> list[CustomerImportRow]:
    matching = next((name for name in workbook.sheetnames if _header_key(name) == _header_key(sheet_name)), None)
    if not matching:
        raise ValueError(f"Không tìm thấy sheet {sheet_name}; hiện có: {', '.join(workbook.sheetnames)}")
    return _parse_rows(workbook[matching].iter_rows(values_only=True))


def import_customers(db: Session, rows: list[CustomerImportRow], *, dry_run: bool = False) -> ImportSummary:
    codes = [row.customer_code for row in rows]
    existing = set(db.scalars(select(Customer.customer_code).where(Customer.customer_code.in_(codes)))) if codes else set()
    imported = 0
    for row in rows:
        if row.customer_code in existing:
            continue
        status_key = _header_key(row.source_status)
        deleted = status_key == "daxoa"
        inactive = deleted or status_key in {"ngungmua", "inactive", "locked"}
        db.add(Customer(
            customer_code=row.customer_code,
            name=row.name,
            phone=row.phone,
            address=row.address,
            customer_group=row.customer_group or "Khách lẻ",
            status="INACTIVE" if inactive else "ACTIVE",
            deleted_at=datetime.now(timezone.utc) if deleted else None,
        ))
        imported += 1
    if dry_run:
        db.flush()
        db.rollback()
    else:
        db.commit()
    return ImportSummary(source_rows=len(rows), imported=imported, skipped_existing=len(existing))


def main() -> None:
    parser = argparse.ArgumentParser(description="Import production tbl_KhachHang vào PostgreSQL")
    parser.add_argument("--mode", choices=("production",), required=True, help="Xác nhận đây là import PRODUCTION")
    parser.add_argument("--file", type=Path, required=True, help="File export 02_DATA_SALES_CRM (.xlsx/.csv/.tsv)")
    parser.add_argument("--sheet", default=SHEET_NAME)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.file.is_file():
        parser.error(f"Không tìm thấy file: {args.file}")
    rows = load_customer_rows(args.file, args.sheet)
    with SessionLocal() as db:
        summary = import_customers(db, rows, dry_run=args.dry_run)
    mode = "DRY-RUN" if args.dry_run else "COMMITTED"
    print(f"{mode}: source={summary.source_rows}, imported={summary.imported}, skipped_existing={summary.skipped_existing}")


if __name__ == "__main__":
    main()
