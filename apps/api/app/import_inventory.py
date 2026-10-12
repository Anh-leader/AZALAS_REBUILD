"""Safe, insert-only import for 03_DATA_KHO_THANH_PHAM exports."""

from __future__ import annotations

import argparse
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .core.db import SessionLocal
from .import_customers import _header_key, _text
from .models import Color, InventoryBalance, Product, ProductVariant, Warehouse, WarehouseLocation

SHEETS = ("tbl_Kho_HangHoa", "tbl_MauSac", "tbl_Kho_BienThe", "tbl_Kho_Ton")

ALIASES = {
    "tbl_Kho_HangHoa": {
        "product_code": ("masp", "mahanghoa"), "pattern_code": ("marap",),
        "brand": ("hieu", "thuonghieu"), "name": ("tensanpham", "tensp", "mausanpham"),
        "category": ("nhomhang", "loaihang"), "status": ("trangthai",), "notes": ("ghichu",),
    },
    "tbl_MauSac": {
        "color_code": ("mamau", "mamausac", "code"), "name": ("tenmau", "mau", "mausac", "ten"),
        "aliases": ("alias", "aliases", "tengoi", "tenthuonggoi"), "status": ("trangthai",),
    },
    "tbl_Kho_BienThe": {
        "internal_code": ("mahangnoibo", "manoibo"), "sales_code": ("mabanhang",),
        "product_code": ("masp",), "color": ("mau", "mauchuan"), "size": ("size", "kichthuoc"),
        "display_name": ("tenhienthi", "tenbienthe"), "status": ("coban", "trangthai"),
    },
    "tbl_Kho_Ton": {
        "internal_code": ("mahangnoibo", "manoibo"), "sales_code": ("mabanhang",),
        "warehouse": ("kho", "khoxuat"), "location": ("vitri",),
        "on_hand": ("tonthucte",), "reserved": ("tongiu",), "available": ("khadung",),
        "pending_finish": ("chohoanthien",), "cut_separated": ("dacotri",), "defective": ("hangloi",),
    },
}

REQUIRED = {
    "tbl_Kho_HangHoa": ("product_code", "name"),
    "tbl_MauSac": ("color_code", "name"),
    "tbl_Kho_BienThe": ("internal_code", "product_code", "size"),
    "tbl_Kho_Ton": ("internal_code", "warehouse", "on_hand", "reserved"),
}


def _records(workbook, sheet_name: str) -> list[tuple[int, dict[str, str | None]]]:
    matching = next((name for name in workbook.sheetnames if _header_key(name) == _header_key(sheet_name)), None)
    if not matching:
        raise ValueError(f"Thiếu sheet bắt buộc: {sheet_name}")
    rows = workbook[matching].iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError(f"{sheet_name} không có header")
    normalized = {_header_key(value): index for index, value in enumerate(headers)}
    columns = {}
    for field, aliases in ALIASES[sheet_name].items():
        index = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if index is not None:
            columns[field] = index
    missing = [field for field in REQUIRED[sheet_name] if field not in columns]
    if missing:
        raise ValueError(f"{sheet_name} thiếu cột: {', '.join(missing)}")
    result = []
    for row_number, row in enumerate(rows, start=2):
        if not any(_text(cell) for cell in row):
            continue
        result.append((row_number, {
            field: _text(row[index]) if index < len(row) else None for field, index in columns.items()
        }))
    return result


def _quantity(value: str | None, *, sheet: str, row: int, field: str) -> Decimal:
    if value is None:
        return Decimal(0)
    try:
        result = Decimal(value.replace(" ", "").replace(",", "."))
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError(f"{sheet} dòng {row}: {field} không phải số: {value}") from exc
    if not result.is_finite() or result.as_tuple().exponent < -3 or result < 0:
        raise ValueError(f"{sheet} dòng {row}: {field} phải >= 0 và tối đa 3 số lẻ")
    return result


def _status(value: str | None) -> str:
    key = _header_key(value)
    return "INACTIVE" if key in {"khong", "ngungban", "inactive", "daxoa"} else "ACTIVE"


def import_inventory(db: Session, workbook, *, dry_run: bool = False) -> dict:
    report = {"dry_run": dry_run, "committed": False, "tables": {}, "errors": []}
    data = {}
    try:
        for sheet in SHEETS:
            data[sheet] = _records(workbook, sheet)
            report["tables"][sheet] = {"source": len(data[sheet]), "imported": 0, "skipped": 0, "errors": 0}
    except ValueError as exc:
        report["errors"].append(str(exc))
        db.rollback()
        return report

    products = {row.product_code: row for row in db.scalars(select(Product))}
    colors_by_code = {row.color_code: row for row in db.scalars(select(Color))}
    colors_by_name = {_header_key(row.canonical_name): row for row in colors_by_code.values()}
    variants = {row.internal_code: row for row in db.scalars(select(ProductVariant))}
    warehouses = {}
    for warehouse in db.scalars(select(Warehouse)):
        warehouses[_header_key(warehouse.code)] = warehouse
        warehouses[_header_key(warehouse.name)] = warehouse
    locations = {}
    for location in db.scalars(select(WarehouseLocation)):
        locations[(location.warehouse_id, _header_key(location.code))] = location
        if location.display_name:
            locations[(location.warehouse_id, _header_key(location.display_name))] = location
    balances = {(row.variant_id, row.warehouse_id, row.location_id): row for row in db.scalars(select(InventoryBalance))}

    def error(sheet: str, row_number: int, text: str) -> None:
        report["tables"][sheet]["errors"] += 1
        report["errors"].append(f"{sheet} dòng {row_number}: {text}")

    for row_number, row in data["tbl_MauSac"]:
        code, name = row.get("color_code"), row.get("name")
        if not code or not name:
            error("tbl_MauSac", row_number, "MaMau và TenMau là bắt buộc; không tự sinh mã")
            continue
        if code in colors_by_code:
            report["tables"]["tbl_MauSac"]["skipped"] += 1
            continue
        color = Color(color_code=code, canonical_name=name, aliases=row.get("aliases"), status=_status(row.get("status")))
        db.add(color); db.flush()
        colors_by_code[code] = color; colors_by_name[_header_key(name)] = color
        report["tables"]["tbl_MauSac"]["imported"] += 1

    for row_number, row in data["tbl_Kho_HangHoa"]:
        code, name = row.get("product_code"), row.get("name")
        if not code or not name:
            error("tbl_Kho_HangHoa", row_number, "MaSP và TenSanPham là bắt buộc; không tự sinh mã")
            continue
        if code in products:
            report["tables"]["tbl_Kho_HangHoa"]["skipped"] += 1
            continue
        product = Product(product_code=code, pattern_code=row.get("pattern_code"), brand=row.get("brand"),
                          name=name, category=row.get("category"), status=_status(row.get("status")), notes=row.get("notes"))
        db.add(product); db.flush(); products[code] = product
        report["tables"]["tbl_Kho_HangHoa"]["imported"] += 1

    dimensions = {(row.product_id, row.color_id, row.size): row for row in variants.values()}
    for row_number, row in data["tbl_Kho_BienThe"]:
        internal, product_code, size = row.get("internal_code"), row.get("product_code"), row.get("size")
        if not internal or not product_code or not size:
            error("tbl_Kho_BienThe", row_number, "MaHangNoiBo, MaSP và Size là bắt buộc")
            continue
        if internal in variants:
            report["tables"]["tbl_Kho_BienThe"]["skipped"] += 1
            continue
        product = products.get(product_code)
        if not product:
            error("tbl_Kho_BienThe", row_number, f"MaSP không tồn tại: {product_code}")
            continue
        color_text = row.get("color")
        color = (colors_by_code.get(color_text) or colors_by_name.get(_header_key(color_text))) if color_text else None
        if color_text and not color:
            error("tbl_Kho_BienThe", row_number, f"Màu không tồn tại trong tbl_MauSac: {color_text}")
            continue
        dimension = (product.id, color.id if color else None, size)
        if dimension in dimensions:
            error("tbl_Kho_BienThe", row_number, f"Trùng product/color/size với mã {dimensions[dimension].internal_code}")
            continue
        display = row.get("display_name")
        if not display:
            error("tbl_Kho_BienThe", row_number, "Thiếu TenHienThi; importer không tự đặt tên")
            continue
        variant = ProductVariant(product_id=product.id, internal_code=internal, sales_code=row.get("sales_code"),
                                 color_id=color.id if color else None, size=size, display_name=display,
                                 status=_status(row.get("status")))
        db.add(variant); db.flush(); variants[internal] = variant; dimensions[dimension] = variant
        report["tables"]["tbl_Kho_BienThe"]["imported"] += 1

    for row_number, row in data["tbl_Kho_Ton"]:
        internal, warehouse_text = row.get("internal_code"), row.get("warehouse")
        variant = variants.get(internal)
        warehouse = warehouses.get(_header_key(warehouse_text))
        if not variant:
            error("tbl_Kho_Ton", row_number, f"MaHangNoiBo không tồn tại: {internal}")
            continue
        if not warehouse:
            error("tbl_Kho_Ton", row_number, f"Kho không khớp chính xác code/name master: {warehouse_text}")
            continue
        location = None
        if row.get("location"):
            location = locations.get((warehouse.id, _header_key(row["location"])))
            if not location:
                # `ViTri` là dữ liệu nguồn trực tiếp, không phải giá trị suy đoán.
                # Vì file nguồn không có sheet master vị trí riêng, lưu nguyên văn
                # làm code/display name và tái sử dụng idempotently cho các dòng sau.
                location = WarehouseLocation(
                    warehouse_id=warehouse.id, code=row["location"], display_name=row["location"]
                )
                db.add(location); db.flush()
                locations[(warehouse.id, _header_key(row["location"]))] = location
        try:
            on_hand = _quantity(row.get("on_hand"), sheet="tbl_Kho_Ton", row=row_number, field="TonThucTe")
            reserved = _quantity(row.get("reserved"), sheet="tbl_Kho_Ton", row=row_number, field="TonGiu")
            pending = _quantity(row.get("pending_finish"), sheet="tbl_Kho_Ton", row=row_number, field="ChoHoanThien")
            cut = _quantity(row.get("cut_separated"), sheet="tbl_Kho_Ton", row=row_number, field="DaCotRi")
            defective = _quantity(row.get("defective"), sheet="tbl_Kho_Ton", row=row_number, field="HangLoi")
        except ValueError as exc:
            error("tbl_Kho_Ton", row_number, str(exc).split(": ", 1)[-1]); continue
        if reserved > on_hand:
            error("tbl_Kho_Ton", row_number, f"TonGiu {reserved} vượt TonThucTe {on_hand}")
            continue
        calculated = on_hand - reserved
        if row.get("available") is not None:
            try:
                source_available = _quantity(row["available"], sheet="tbl_Kho_Ton", row=row_number, field="KhaDung")
            except ValueError as exc:
                error("tbl_Kho_Ton", row_number, str(exc).split(": ", 1)[-1]); continue
            if source_available != calculated:
                error("tbl_Kho_Ton", row_number, f"KhaDung nguồn {source_available} != TonThucTe-TonGiu {calculated}")
                continue
        scope = (variant.id, warehouse.id, location.id if location else None)
        if scope in balances:
            report["tables"]["tbl_Kho_Ton"]["skipped"] += 1
            continue
        balance = InventoryBalance(variant_id=variant.id, warehouse_id=warehouse.id,
                                   location_id=location.id if location else None, on_hand=on_hand,
                                   reserved=reserved, available=calculated, pending_finish=pending,
                                   cut_separated=cut, defective=defective)
        db.add(balance); db.flush(); balances[scope] = balance
        report["tables"]["tbl_Kho_Ton"]["imported"] += 1

    if report["errors"] or dry_run:
        db.rollback()
    else:
        db.commit(); report["committed"] = True
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Import 03_DATA_KHO_THANH_PHAM an toàn, insert-only")
    parser.add_argument("--mode", choices=("production",), required=True)
    parser.add_argument("--file", type=Path, required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.file.is_file():
        parser.error(f"Không tìm thấy file: {args.file}")
    workbook = load_workbook(args.file, read_only=True, data_only=True)
    try:
        with SessionLocal() as db:
            report = import_inventory(db, workbook, dry_run=args.dry_run)
    finally:
        workbook.close()
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    if report["errors"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
