from openpyxl import Workbook
from sqlalchemy import select

from app.import_inventory import import_inventory
from app.models import InventoryBalance, Product, ProductVariant, WarehouseLocation


def _workbook(warehouse="Tân Xuân"):
    workbook = Workbook()
    products = workbook.active
    products.title = "tbl_Kho_HangHoa"
    products.append(["MaSP", "MaRap", "Hieu", "TenSanPham", "NhomHang", "TrangThai"])
    products.append(["SP-SOURCE-1", "RAP-01", "AZALAS", "Áo nguồn", "Áo", "Đang bán"])
    colors = workbook.create_sheet("tbl_MauSac")
    colors.append(["MaMau", "TenMau", "Alias", "TrangThai"])
    colors.append(["MAU-XANH", "Xanh", "xanh bien", "ACTIVE"])
    variants = workbook.create_sheet("tbl_Kho_BienThe")
    variants.append(["MaHangNoiBo", "MaBanHang", "MaSP", "Mau", "Size", "TenHienThi", "CoBan"])
    variants.append(["NB-SOURCE-1", "BH-SOURCE-1", "SP-SOURCE-1", "Xanh", "L", "Áo nguồn Xanh L", "Có"])
    stock = workbook.create_sheet("tbl_Kho_Ton")
    stock.append(["MaHangNoiBo", "Kho", "ViTri", "TonThucTe", "TonGiu", "KhaDung", "ChoHoanThien", "DaCotRi", "HangLoi"])
    stock.append(["NB-SOURCE-1", warehouse, None, 10, 2, 8, 1, 0, 0])
    return workbook


def test_inventory_import_dry_run_rolls_back(db):
    workbook = _workbook()
    report = import_inventory(db, workbook, dry_run=True)
    workbook.close()
    assert not report["errors"] and not report["committed"]
    assert report["tables"]["tbl_Kho_Ton"]["imported"] == 1
    assert db.scalar(select(Product).where(Product.product_code == "SP-SOURCE-1")) is None


def test_inventory_import_is_idempotent_and_insert_only(db):
    workbook = _workbook()
    first = import_inventory(db, workbook)
    second = import_inventory(db, workbook)
    workbook.close()
    assert first["committed"] and not first["errors"]
    assert second["committed"] and not second["errors"]
    assert second["tables"]["tbl_Kho_HangHoa"]["skipped"] == 1
    assert second["tables"]["tbl_Kho_BienThe"]["skipped"] == 1
    assert second["tables"]["tbl_Kho_Ton"]["skipped"] == 1
    variant = db.scalar(select(ProductVariant).where(ProductVariant.internal_code == "NB-SOURCE-1"))
    balances = list(db.scalars(select(InventoryBalance).where(InventoryBalance.variant_id == variant.id)))
    assert len(balances) == 1 and balances[0].on_hand == 10 and balances[0].reserved == 2


def test_inventory_import_unknown_warehouse_rolls_back(db):
    workbook = _workbook("Kho tự suy đoán")
    report = import_inventory(db, workbook)
    workbook.close()
    assert not report["committed"]
    assert any("Kho không khớp chính xác" in error for error in report["errors"])
    assert db.scalar(select(Product).where(Product.product_code == "SP-SOURCE-1")) is None


def test_inventory_import_preserves_source_location_without_guessing(db):
    workbook = _workbook()
    workbook["tbl_Kho_Ton"]["C2"] = "Kệ A-01"
    report = import_inventory(db, workbook)
    workbook.close()
    assert report["committed"] and not report["errors"]
    location = db.scalar(select(WarehouseLocation).where(WarehouseLocation.code == "Kệ A-01"))
    assert location is not None and location.display_name == "Kệ A-01"
    balance = db.scalar(select(InventoryBalance).where(InventoryBalance.location_id == location.id))
    assert balance is not None
