from pathlib import Path

from sqlalchemy import select

from app.import_customers import CustomerImportRow, _load_workbook_rows, import_customers, load_customer_rows
from app.models import Customer, InventoryBalance, Product, Warehouse
from app.seed import LEGACY_DEMO_CUSTOMER_CODES, demo_customer_generation, seed_demo_inventory


def test_demo_customers_never_have_generated_phone(db):
    customers = list(db.scalars(select(Customer)))
    assert customers
    assert all(customer.customer_code.startswith("DEMO-KH-") for customer in customers)
    assert all(customer.phone is None for customer in customers)


def test_demo_detection_never_guesses_production_customers():
    assert demo_customer_generation(set(LEGACY_DEMO_CUSTOMER_CODES)) == "LEGACY"
    assert demo_customer_generation({"DEMO-KH-0001"}) == "CURRENT"
    assert demo_customer_generation({"KH-SOURCE-001"}) is None
    assert demo_customer_generation(set(LEGACY_DEMO_CUSTOMER_CODES) | {"KH-REAL"}) is None


def test_demo_inventory_backfill_is_idempotent(db):
    warehouses = {row.code: row for row in db.scalars(select(Warehouse))}
    before = (db.query(Product).count(), db.query(InventoryBalance).count())
    seed_demo_inventory(db, (warehouses["TAN_XUAN"], warehouses["CA_KOI"]))
    seed_demo_inventory(db, (warehouses["TAN_XUAN"], warehouses["CA_KOI"]))
    db.commit()
    assert (db.query(Product).count(), db.query(InventoryBalance).count()) == before


def test_import_reads_phone_only_when_sdt_has_value_and_never_overwrites(db):
    existing = db.scalar(select(Customer).where(Customer.customer_code == "DEMO-KH-0001"))
    original_name = existing.name
    source = Path(__file__).parent / "fixtures" / "tbl_KhachHang.csv"
    rows = load_customer_rows(source)
    assert rows[1].phone is None
    assert rows[2].phone == "0912345678"
    result = import_customers(db, rows)

    assert result.imported == 2
    assert result.skipped_existing == 1
    db.refresh(existing)
    assert existing.name == original_name
    assert existing.phone is None
    assert db.scalar(select(Customer).where(Customer.customer_code == "KH-SOURCE-001")).phone is None
    assert db.scalar(select(Customer).where(Customer.customer_code == "KH-SOURCE-002")).phone == "0912345678"


def test_import_dry_run_rolls_back(db):
    rows = [CustomerImportRow(2, "KH-DRY-RUN", "Khách dry run", None, None, None, None)]
    result = import_customers(db, rows, dry_run=True)
    assert result.imported == 1
    assert db.scalar(select(Customer).where(Customer.customer_code == "KH-DRY-RUN")) is None


def test_import_xlsx_uses_tbl_khach_hang_sheet():
    from openpyxl import Workbook

    workbook = Workbook()
    other = workbook.active
    other.title = "tbl_DonHang"
    sheet = workbook.create_sheet("tbl_KhachHang")
    sheet.append(["MaKH", "TenKhach", "SDT", "DiaChi"])
    sheet.append(["KH-XLSX-001", "Khách XLSX", None, "Huế"])
    rows = _load_workbook_rows(workbook)
    workbook.close()
    assert len(rows) == 1
    assert rows[0].customer_code == "KH-XLSX-001"
    assert rows[0].phone is None
